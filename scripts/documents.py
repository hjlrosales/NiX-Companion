"""Local document worker. JSON in/out; no network or shell interpolation."""
import json, os, sys, re, shutil, subprocess, tempfile, collections, zipfile
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from pptx.dml.color import RGBColor as SlideColor

def safe_archive(path):
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100_000_000 or len(archive.infolist()) > 10000:
            raise ValueError('Document archive exceeds the extraction limit')

def extract(path):
    if path.stat().st_size > 50_000_000: raise ValueError('Document exceeds 50 MB')
    ext = path.suffix.lower()
    parts = []
    if ext == '.pdf':
        with fitz.open(path) as doc:
            if len(doc) > 300: raise ValueError('PDF exceeds 300 pages')
            for i, page in enumerate(doc):
                text = page.get_text()
                ocr = False
                if len(text.strip()) < 30:
                    tessdata = os.environ.get('TESSDATA_PREFIX', r'C:\Program Files\Tesseract-OCR\tessdata')
                    if not Path(tessdata, 'eng.traineddata').exists(): raise ValueError(f'Page {i+1} requires local OCR. Install Tesseract with English language data first.')
                    text = page.get_text(textpage=page.get_textpage_ocr(language='eng', dpi=150, full=True, tessdata=tessdata)); ocr=True
                parts.append({'reference': f'{path.name}#page={i+1}', 'text': text, 'ocr': ocr})
    elif ext == '.docx':
        safe_archive(path); doc=Document(path)
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip(): parts.append({'reference':f'{path.name}#paragraph={i+1}','text':p.text})
        for i, table in enumerate(doc.tables):
            parts.append({'reference':f'{path.name}#table={i+1}','text':'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)})
    elif ext == '.pptx':
        safe_archive(path); deck=Presentation(path)
        for i, slide in enumerate(deck.slides):
            content=[]
            for shape in slide.shapes:
                if shape.has_text_frame: content.append(shape.text)
                if shape.has_table: content.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            parts.append({'reference':f'{path.name}#slide={i+1}','text':'\n'.join(content)})
    elif ext in ['.txt','.md','.csv','.log','.json']:
        text=path.read_text(encoding='utf-8')
        lines=text.splitlines()
        for start in range(0,len(lines),40): parts.append({'reference':f'{path.name}#line={start+1}','text':'\n'.join(lines[start:start+40])})
    elif ext in ['.png','.jpg','.jpeg','.tiff','.bmp']:
        with fitz.open(path) as image:
            with fitz.open('pdf',image.convert_to_pdf()) as pdf:
                page=pdf[0]; tessdata=os.environ.get('TESSDATA_PREFIX',r'C:\Program Files\Tesseract-OCR\tessdata')
                text=page.get_text(textpage=page.get_textpage_ocr(language='eng', dpi=150, full=True, tessdata=tessdata))
                parts=[{'reference':f'{path.name}#image=1','text':text,'ocr':True}]
    else: raise ValueError('Supported inputs: PDF, DOCX, PPTX, text, Markdown, CSV, JSON and scanned images')
    if sum(len(p['text']) for p in parts)>2_000_000: raise ValueError('Extracted text exceeds 2 million characters')
    return parts

def summary(parts):
    """Deterministic extractive summary: retain source sentences and exact references."""
    result=[]
    for part in parts:
        lines=part['text'].splitlines()
        first_line=lines[0].strip() if lines else ''
        prose=' '.join(lines[1:] if 3<=len(first_line)<=90 and len(lines)>1 else lines)
        prose=re.sub(r'\s+',' ',prose).strip()
        sentences=list(dict.fromkeys(s.strip() for s in re.split(r'(?<=[.!?])\s+',prose) if len(s.strip())>=30))
        if not sentences: sentences=[part['text'].strip()]
        words=collections.Counter(re.findall(r'\b[a-z]{4,}\b',part['text'].lower()))
        scored=sorted(enumerate(sentences[1:],1),key=lambda x:sum(words[w] for w in re.findall(r'\b[a-z]{4,}\b',x[1].lower()))/max(1,len(x[1])**.5),reverse=True)[:2]
        text=' '.join(sentences[i] for i in sorted({0,*[i for i,_ in scored]}))
        first_line=part['text'].splitlines()[0].strip() if part['text'].splitlines() else ''
        heading=first_line if 3<=len(first_line)<=90 else f'Source section {len(result)+1}'
        if text: result.append({'title':heading,'text':text,'reference':part['reference']})
    return result

def office_pdf(path,out):
    pdf=out/(path.stem+'.pdf')
    if os.name=='nt' and not os.environ.get('NIX_SOFFICE'):
        result=subprocess.run(['powershell.exe','-NoLogo','-NoProfile','-NonInteractive','-File',str(Path(__file__).with_name('render-office.ps1')),'-InputPath',str(path.resolve()),'-OutputPath',str(pdf.resolve())],capture_output=True,text=True,timeout=120)
        if result.returncode or not pdf.exists(): raise ValueError('Microsoft Office rendering failed. Verify Word and PowerPoint are installed and activated. '+result.stderr[-2000:])
        return pdf
    office=os.environ.get('NIX_SOFFICE') or shutil.which('soffice') or r'C:\Program Files\LibreOffice\program\soffice.exe'
    if not Path(office).exists(): raise ValueError('Rendering requires Microsoft Office on Windows, or an explicitly configured NIX_SOFFICE executable.')
    with tempfile.TemporaryDirectory(prefix='nix-office-') as profile:
        result=subprocess.run([office,f'-env:UserInstallation={Path(profile).as_uri()}','--headless','--convert-to','pdf','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=90)
        pdf=out/(path.stem+'.pdf')
        if result.returncode or not pdf.exists(): raise ValueError('Office rendering failed: '+result.stdout+result.stderr)
    return pdf

def previews(pdf,out):
    paths=[]
    with fitz.open(pdf) as doc:
        for i,page in enumerate(doc):
            target=out/f'{pdf.stem}-page-{i+1}.png'; page.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(target); paths.append(str(target))
    return paths

def report(title,sections,out):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8)
    style=doc.styles['Normal']; style.font.name='Calibri'; style.font.size=Pt(11); style.paragraph_format.space_after=Pt(8)
    for name in ['Title','Heading 1','Heading 2','Caption']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        for border in list(doc.styles[name].element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.add_heading(title,0)
    doc.add_paragraph('This report presents source-linked findings. Extractive summaries preserve selected source sentences; consult the referenced sections for full context.')
    for s in sections:
        doc.add_heading(s['title'],1); doc.add_paragraph(s['text']); p=doc.add_paragraph('Source: '+s.get('reference','User supplied content')); p.style='Caption'
    target=out/'report.docx'; doc.save(target)
    reopened=Document(target)
    if len(reopened.paragraphs)<len(sections)*3: raise ValueError('Report content verification failed')
    pdf=office_pdf(target,out)
    return [str(target),str(pdf)],previews(pdf,out)

def slides(title,sections,out):
    deck=Presentation(); deck.slide_width=SlideInches(13.333); deck.slide_height=SlideInches(7.5)
    # Text is editable and deliberately bounded; overflow becomes additional slides.
    pages=[]
    for s in sections:
        text=s['text']; chunks=[]
        while text:
            split=min(len(text),650)
            if split<len(text): split=max(text.rfind(' ',0,split),1)
            chunks.append(text[:split]); text=text[split:].lstrip()
        for n,chunk in enumerate(chunks): pages.append((s['title']+(f' continued {n+1}' if n else ''),chunk,s.get('reference','User supplied content')))
    for index,(heading,body,ref) in enumerate([(title,'Source-linked briefing','Overview')]+pages):
        slide=deck.slides.add_slide(deck.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=SlideColor(18,26,34)
        def box(x,y,w,h,text,size,color):
            shape=slide.shapes.add_textbox(SlideInches(x),SlideInches(y),SlideInches(w),SlideInches(h)); tf=shape.text_frame; tf.word_wrap=True
            for line_number,line in enumerate(text.split('\n')):
                p=tf.paragraphs[0] if line_number==0 else tf.add_paragraph(); p.text=line; p.font.name='Calibri'; p.font.size=SlidePt(size); p.font.color.rgb=SlideColor(*color); p.space_after=SlidePt(12)
            return shape
        box(.7,.42,11.7,.45,f'NIX  /  {index+1:02d}',11,(133,194,212))
        box(.7,1.15,11.7,1.0,heading[:110],30,(232,240,246))
        box(.7,2.5,11.6,3.65,body,22,(192,210,223))
        box(.7,6.75,11.7,.4,ref[:170],11,(135,163,181))
    target=out/'briefing.pptx'; deck.save(target); reopened=Presentation(target)
    if len(reopened.slides)!=len(pages)+1: raise ValueError('Slide count verification failed')
    pdf=office_pdf(target,out)
    return [str(target),str(pdf)],previews(pdf,out)

def main(request):
    action=request['action']; source=Path(request['source']) if request.get('source') else None
    if action=='extract':
        parts=extract(source); start=request.get('start',0); selected=parts[start:start+request.get('count',3)]
        return {'output':json.dumps({'totalSections':len(parts),'sections':selected},ensure_ascii=False)[:12000], 'evidence':[f'Extracted {len(parts)} source sections from {source.name}']}
    out=Path(request['output']); out.mkdir(parents=True,exist_ok=False)
    sections=request.get('sections')
    if source:
        parts=extract(source); sections=summary(parts)
        if not sections: raise ValueError('No readable source content was found')
        (out/'extraction.json').write_text(json.dumps(parts,ensure_ascii=False,indent=2),encoding='utf-8')
    if not sections: raise ValueError('Supply source content or sections')
    if len(sections)>300: raise ValueError('Deliverable exceeds 300 sections')
    summary_file=out/'summary.md'; summary_file.write_text('# '+request.get('title','Source summary')+'\n\nExtractive summary with traceable source references.\n\n'+'\n\n'.join(f"## {s['title']}\n\n{s['text']}\n\nSource: {s.get('reference','User supplied content')}" for s in sections),encoding='utf-8')
    artifacts=[str(summary_file)]; images=[]
    if action in ['report','bundle']:
        files,rendered=report(request.get('title','Source report'),sections,out); artifacts+=files; images+=rendered
    if action in ['slides','bundle']:
        files,rendered=slides(request.get('title','Source briefing'),sections,out); artifacts+=files; images+=rendered
    return {'output':f'Created {len(artifacts)} files from {len(sections)} source-linked sections. Office files reopened and rendered. Inspect the rendered pages before accepting.','artifacts':artifacts+images,'evidence':[f'{len(sections)} source-linked sections',f'{len(images)} rendered pages/slides'],'previews':images}

if __name__=='__main__':
    try: print(json.dumps(main(json.load(sys.stdin)),ensure_ascii=False))
    except Exception as error: print(json.dumps({'error':str(error)})); sys.exit(1)
