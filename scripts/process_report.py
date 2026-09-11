"""Create a DOCX process report from JSON stdin."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt


def main(request):
    output = Path(request["output"])
    processes = request["processes"]
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)
    doc.add_heading("Running Processes Report", 0)
    doc.add_paragraph(
        "This report was generated from the current Windows process list. "
        "Some fields, such as path, company, or CPU time, may be blank when Windows does not expose that information to the current user."
    )
    doc.add_paragraph(f"Processes captured: {len(processes)}")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Process", "PID", "CPU", "Memory MB", "Company", "Plain-English explanation"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for item in processes:
        row = table.add_row().cells
        row[0].text = str(item.get("name") or "")
        row[1].text = str(item.get("id") or "")
        row[2].text = "" if item.get("cpu") is None else str(item.get("cpu"))
        row[3].text = "" if item.get("memoryMb") is None else str(item.get("memoryMb"))
        row[4].text = str(item.get("company") or "")
        row[5].text = str(item.get("explanation") or "")
    doc.add_heading("Notes", 1)
    doc.add_paragraph("High memory use can be normal for browsers, design tools, local AI runtimes, and Electron applications.")
    doc.add_paragraph("Processes without a visible path or company are often protected system services or processes owned by another account.")
    doc.save(output)
    reopened = Document(output)
    if len(reopened.tables) != 1 or len(reopened.tables[0].rows) < 2:
        raise ValueError("DOCX verification failed after saving.")
    return {"output": str(output), "count": len(processes)}


if __name__ == "__main__":
    try:
        print(json.dumps(main(json.load(sys.stdin))))
    except Exception as error:
        print(json.dumps({"error": str(error)}))
        sys.exit(1)
